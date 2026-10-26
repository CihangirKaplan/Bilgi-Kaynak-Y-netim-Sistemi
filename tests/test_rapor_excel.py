"""Cihaz raporu Excel çıktısı testleri (Öğrenci B - DASH-006)."""
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from rapor_excel import excel_dosyasi_olustur

BUGUN = date(2026, 10, 26)

VERI = {
    "ozet": {
        "bakim_sayisi": 3,
        "toplam_bakim_maliyeti": 1800.0,
        "kalibrasyon_sayisi": 1,
        "ariza_sayisi": 3,
        "acik_ariza_sayisi": 1,
    },
    "cihazlar": [
        {"cihaz_id": 1, "envanter_kodu": "CHZ-1", "cihaz_adi": "EEG", "marka_model": "Marka X",
         "durum": "Bakimda", "sonraki_bakim_tarihi": "2027-04-01",
         "son_kalibrasyon": {"kalibrasyon_yapan_kurum": "Kurum A", "sertifika_no": "S-1",
                             "islem_tarihi": "2026-10-02T09:00:00", "gecerlilik_tarihi": "2026-11-10"}},
        {"cihaz_id": 2, "envanter_kodu": "CHZ-2", "cihaz_adi": "EKG", "marka_model": None,
         "durum": "Arizali", "sonraki_bakim_tarihi": "2026-10-20",
         "son_kalibrasyon": {"kalibrasyon_yapan_kurum": "Kurum B", "sertifika_no": "S-2",
                             "islem_tarihi": "2025-01-02T09:00:00", "gecerlilik_tarihi": "2026-01-02"}},
        {"cihaz_id": 3, "envanter_kodu": "CHZ-3", "cihaz_adi": "Monitör", "marka_model": None,
         "durum": "Kullanima_Hazir", "sonraki_bakim_tarihi": None,
         "son_kalibrasyon": {"kalibrasyon_yapan_kurum": "Kurum C", "sertifika_no": "S-3",
                             "islem_tarihi": "2026-09-01T09:00:00", "gecerlilik_tarihi": "2027-09-01"}},
        {"cihaz_id": 4, "envanter_kodu": "CHZ-4", "cihaz_adi": "Pulse", "marka_model": None,
         "durum": "Kullanima_Hazir", "sonraki_bakim_tarihi": None, "son_kalibrasyon": None},
    ],
    "bakimlar": [
        # devam ediyor (bugün başlangıç ile bitiş arasında)
        {"cihaz_id": 1, "cihaz": "CHZ-1 - EEG", "bakim_tarihi": "2026-10-24",
         "bakim_baslangic_tarihi": "2026-10-24", "bakim_bitis_tarihi": "2026-10-30",
         "bakim_yapan_kisi": "Ali", "yapilan_islem": "Filtre değişimi",
         "maliyet": 1000.0, "sonraki_bakim_tarihi": "2027-04-01"},
        # tamamlandı
        {"cihaz_id": 1, "cihaz": "CHZ-1 - EEG", "bakim_tarihi": "2026-09-01",
         "bakim_baslangic_tarihi": "2026-09-01", "bakim_bitis_tarihi": "2026-09-03",
         "bakim_yapan_kisi": "Veli", "yapilan_islem": "Genel bakım",
         "maliyet": 500.0, "sonraki_bakim_tarihi": None},
        # planlanmış
        {"cihaz_id": 3, "cihaz": "CHZ-3 - Monitör", "bakim_tarihi": "2026-11-05",
         "bakim_baslangic_tarihi": "2026-11-05", "bakim_bitis_tarihi": "2026-11-06",
         "bakim_yapan_kisi": "Can", "yapilan_islem": "Pil değişimi",
         "maliyet": 300.0, "sonraki_bakim_tarihi": None},
    ],
    "kalibrasyonlar": [
        {"cihaz_id": 1, "cihaz": "CHZ-1 - EEG", "islem_tarihi": "2026-10-02T09:00:00",
         "kalibrasyon_yapan_kurum": "Kurum A", "sertifika_no": "S-1",
         "gecerlilik_tarihi": "2026-11-10"},
    ],
    "arizalar": [
        {"cihaz_id": 2, "cihaz": "CHZ-2 - EKG", "bildirim_tarihi": "2026-10-20T09:00:00",
         "bildiren_personel": "Ayşe", "ariza_aciklamasi": "Ekran gitti",
         "durum": "Açık", "cozum_tarihi": None},
        {"cihaz_id": 1, "cihaz": "CHZ-1 - EEG", "bildirim_tarihi": "2026-10-04T09:00:00",
         "bildiren_personel": None, "ariza_aciklamasi": "Kablo kopuk",
         "durum": "Çözüldü", "cozum_tarihi": "2026-10-06T09:00:00"},
        {"cihaz_id": 3, "cihaz": "CHZ-3 - Monitör", "bildirim_tarihi": "2026-10-01T09:00:00",
         "bildiren_personel": "Mert", "ariza_aciklamasi": "Açılmıyor",
         "durum": "Çözüldü", "cozum_tarihi": "2026-10-03T09:00:00"},
    ],
}
FILTRE = [("Dönem", "Tümü"), ("Cihaz", "Tüm Cihazlar")]


def olustur(tmp_path, veri=VERI, filtre=FILTRE):
    yol = tmp_path / "rapor.xlsx"
    excel_dosyasi_olustur(veri, filtre, str(yol), bugun=BUGUN)
    return load_workbook(yol)


def satirlar(sayfa):
    """Özet sayfasını {etiket: (değer, açıklama)} sözlüğüne çevirir."""
    return {
        sayfa.cell(r, 1).value: (sayfa.cell(r, 2).value, sayfa.cell(r, 3).value)
        for r in range(1, sayfa.max_row + 1)
        if sayfa.cell(r, 1).value
    }


def tablolar(sayfa):
    """Durum Raporu'nu {tablo başlığı: [satır sözlükleri]} biçimine çevirir."""
    sonuc = {}
    r = 1
    while r <= sayfa.max_row:
        deger = sayfa.cell(r, 1).value
        if isinstance(deger, str) and " kayıt)" in deger:
            baslik = deger.split("  (")[0]
            r += 2  # bölüm başlığı + açıklama
            kayitlar = []
            if sayfa.cell(r, 1).value == "Bu kategoride kayıt yok.":
                r += 1
            else:
                kolonlar = [sayfa.cell(r, c).value for c in range(1, 8)]
                r += 1
                while sayfa.cell(r, 1).value:
                    kayitlar.append({kolonlar[c - 1]: sayfa.cell(r, c).value for c in range(1, 8)})
                    r += 1
            sonuc[baslik] = kayitlar
        else:
            r += 1
    return sonuc


def test_alti_sayfa_olusur(tmp_path):
    kitap = olustur(tmp_path)
    assert kitap.sheetnames == [
        "Özet", "Durum Raporu", "Cihaz Özeti", "Bakım Geçmişi", "Kalibrasyon", "Arızalar"
    ]


def test_filtreler_aciklamalariyla_yazilir(tmp_path):
    ozet = satirlar(olustur(tmp_path)["Özet"])

    assert ozet["Dönem"][0] == "Tümü"
    assert "kayıtlı tüm dönemler" in ozet["Dönem"][1]
    assert ozet["Cihaz"][0] == "Tüm Cihazlar"
    assert "tüm cihazlar rapora dahil" in ozet["Cihaz"][1]
    assert ozet["Rapordaki cihazlar"][0] == 4
    assert "CHZ-4 (Pulse)" in ozet["Rapordaki cihazlar"][1]


def test_secili_donem_ve_cihaz_aciklamasi(tmp_path):
    filtre = [("Dönem", "01.10.2026 - 20.10.2026"), ("Cihaz", "CHZ-1 - EEG")]
    ozet = satirlar(olustur(tmp_path, filtre=filtre)["Özet"])

    assert "Bu tarihler arasındaki" in ozet["Dönem"][1]
    assert "yalnızca bu cihaza" in ozet["Cihaz"][1]


def test_bakim_gostergeleri_durumlara_ayrilir(tmp_path):
    ozet = satirlar(olustur(tmp_path)["Özet"])

    assert ozet["Bakım kaydı (toplam)"][0] == 3
    assert ozet["Devam eden bakım"][0] == 1
    assert "CHZ-1 - EEG" in ozet["Devam eden bakım"][1]
    assert ozet["Planlanmış bakım"][0] == 1
    assert "CHZ-3 - Monitör" in ozet["Planlanmış bakım"][1]
    assert ozet["Tamamlanan bakım"][0] == 1
    assert ozet["Toplam bakım maliyeti"][0] == 1800.0
    assert ozet["Ortalama bakım maliyeti"][0] == 600.0
    assert ozet["En yüksek maliyetli bakım"][0] == 1000.0
    assert "Filtre değişimi" in ozet["En yüksek maliyetli bakım"][1]


def test_kalibrasyon_gostergeleri_uc_kategoriye_ayrilir(tmp_path):
    ozet = satirlar(olustur(tmp_path)["Özet"])

    assert ozet["Süresi dolmuş"][0] == 1
    assert "CHZ-2" in ozet["Süresi dolmuş"][1]
    assert ozet["30 gün içinde dolacak"][0] == 1
    assert "CHZ-1" in ozet["30 gün içinde dolacak"][1]
    assert ozet["30 günden fazla süresi var"][0] == 1
    assert "CHZ-3" in ozet["30 günden fazla süresi var"][1]
    assert ozet["Kalibrasyon kaydı yok"][0] == 1
    assert "CHZ-4" in ozet["Kalibrasyon kaydı yok"][1]


def test_ariza_gostergeleri(tmp_path):
    ozet = satirlar(olustur(tmp_path)["Özet"])

    assert ozet["Çözümü süren arıza"][0] == 1
    assert "CHZ-2 - EKG" in ozet["Çözümü süren arıza"][1]
    assert ozet["Çözülen arıza"][0] == 2
    assert ozet["Ortalama çözüm süresi"][0] == 2.0
    assert ozet["En uzun süredir açık arıza"][0] == 6


def test_durum_dagilimi_hangi_cihazlar_bakimda_yazar(tmp_path):
    ozet = satirlar(olustur(tmp_path)["Özet"])

    assert ozet["Bakımda"][0] == 1
    assert "CHZ-1 (EEG)" in ozet["Bakımda"][1]
    assert "CHZ-2 (EKG)" in ozet["Arızalı"][1]
    assert ozet["Kullanıma Hazır"][0] == 2


def test_durum_raporu_bakim_tablolari(tmp_path):
    t = tablolar(olustur(tmp_path)["Durum Raporu"])

    devam = t["BAKIM — Devam eden"]
    assert [r["Cihaz"] for r in devam] == ["CHZ-1 - EEG"]
    assert devam[0]["Kaç Gündür Sürüyor"] == 2
    assert devam[0]["Yapılan İşlem"] == "Filtre değişimi"
    assert devam[0]["Maliyet"] == 1000.0

    planli = t["BAKIM — Planlanmış"]
    assert [r["Cihaz"] for r in planli] == ["CHZ-3 - Monitör"]
    assert planli[0]["Başlamasına Kalan Gün"] == 10

    biten = t["BAKIM — Tamamlanan"]
    assert [r["Cihaz"] for r in biten] == ["CHZ-1 - EEG"]
    assert biten[0]["Süre (gün)"] == 2


def test_durum_raporu_sonraki_planli_bakim(tmp_path):
    kayitlar = tablolar(olustur(tmp_path)["Durum Raporu"])["BAKIM — Bir sonraki planlı bakım"]

    # en acil (geciken) üstte
    assert [r["Cihaz"] for r in kayitlar] == ["CHZ-2 - EKG", "CHZ-1 - EEG"]
    assert kayitlar[0]["Kalan Gün"] == -6
    assert kayitlar[0]["Durum"] == "6 gün gecikmiş"
    assert kayitlar[1]["Durum"] == "Zamanı var"


def test_durum_raporu_ariza_tablolari(tmp_path):
    t = tablolar(olustur(tmp_path)["Durum Raporu"])

    acik = t["ARIZA — Çözümü süren"]
    assert [r["Cihaz"] for r in acik] == ["CHZ-2 - EKG"]
    assert acik[0]["Kaç Gündür Açık"] == 6
    assert acik[0]["Arıza Açıklaması"] == "Ekran gitti"

    cozulen = t["ARIZA — Çözülen"]
    assert [r["Cihaz"] for r in cozulen] == ["CHZ-1 - EEG", "CHZ-3 - Monitör"]
    assert cozulen[0]["Çözüm Süresi (gün)"] == 2
    assert cozulen[0]["Bildiren"] is None


def test_durum_raporu_kalibrasyon_tablolari(tmp_path):
    t = tablolar(olustur(tmp_path)["Durum Raporu"])

    dolmus = t["KALİBRASYON — Süresi dolmuş"]
    assert [r["Cihaz"] for r in dolmus] == ["CHZ-2 - EKG"]
    assert dolmus[0]["Kaç Gün Önce Doldu"] == 297
    assert dolmus[0]["Sertifika No"] == "S-2"
    assert dolmus[0]["Kalibrasyonu Yapan Kurum"] == "Kurum B"

    yakinda = t["KALİBRASYON — 30 gün içinde dolacak"]
    assert [r["Cihaz"] for r in yakinda] == ["CHZ-1 - EEG"]
    assert yakinda[0]["Kalan Gün"] == 15

    rahat = t["KALİBRASYON — 30 günden fazla süresi var"]
    assert [r["Cihaz"] for r in rahat] == ["CHZ-3 - Monitör"]
    assert rahat[0]["Kalan Gün"] == 310

    yok = t["KALİBRASYON — Kaydı olmayan"]
    assert [r["Cihaz"] for r in yok] == ["CHZ-4 - Pulse"]


def test_kalibrasyon_durumu_filtreden_bagimsiz_cihaz_bilgisinden_gelir(tmp_path):
    """Dönemde kalibrasyon kaydı olmasa bile cihazın son kalibrasyonu değerlendirilir."""
    veri = {**VERI, "kalibrasyonlar": []}
    t = tablolar(olustur(tmp_path, veri)["Durum Raporu"])

    assert [r["Cihaz"] for r in t["KALİBRASYON — Süresi dolmuş"]] == ["CHZ-2 - EKG"]


def test_cihaz_ozeti_her_cihaz_icin_bir_satir(tmp_path):
    sayfa = olustur(tmp_path)["Cihaz Özeti"]
    baslik = [c.value for c in sayfa[1]]
    satir = {
        sayfa.cell(r, 1).value: {baslik[i]: sayfa.cell(r, i + 1).value for i in range(len(baslik))}
        for r in range(2, sayfa.max_row + 1)
    }

    assert list(satir) == ["CHZ-1", "CHZ-2", "CHZ-3", "CHZ-4"]
    eeg = satir["CHZ-1"]
    assert eeg["Güncel Durum"] == "Bakımda"
    assert eeg["Bakım Sayısı (dönemde)"] == 2
    assert eeg["Devam Eden Bakım"] == 1
    assert eeg["Toplam Bakım Maliyeti"] == 1500.0
    assert eeg["Ortalama Bakım Maliyeti"] == 750.0
    assert eeg["Son Bakım İşlemi"] == "Filtre değişimi"
    assert eeg["Kalibrasyon Durumu"] == "30 gün içinde dolacak"
    assert satir["CHZ-2"]["Çözümü Süren Arıza"] == 1
    assert satir["CHZ-2"]["Son / Açık Arıza Açıklaması"] == "Ekran gitti"
    assert satir["CHZ-2"]["Kalibrasyon Durumu"] == "Süresi dolmuş"
    assert satir["CHZ-3"]["Kalibrasyon Durumu"] == "30 günden fazla süresi var"
    assert satir["CHZ-4"]["Bakım Sayısı (dönemde)"] == 0
    assert satir["CHZ-4"]["Ortalama Bakım Maliyeti"] is None
    assert satir["CHZ-4"]["Kalibrasyon Durumu"] == "Kalibrasyon kaydı yok"


def test_bakim_gecmisi_durum_ve_tarihler(tmp_path):
    sayfa = olustur(tmp_path)["Bakım Geçmişi"]

    assert [sayfa.cell(r, 4).value for r in (2, 3, 4)] == [
        "Devam ediyor", "Tamamlandı", "Planlanmış (henüz başlamadı)"
    ]
    assert sayfa["B2"].value == datetime(2026, 10, 24)
    assert sayfa["C2"].value == datetime(2026, 10, 30)
    assert sayfa["G2"].value == 1000.0


def test_kalibrasyon_ve_ariza_sayfalarindaki_ek_sutunlar(tmp_path):
    kitap = olustur(tmp_path)

    assert kitap["Kalibrasyon"]["F2"].value == "30 gün içinde dolacak"
    ariza = kitap["Arızalar"]
    assert [ariza.cell(r, 5).value for r in (2, 3)] == ["Çözümü sürüyor", "Çözüldü"]
    assert ariza["G2"].value == 6      # 20.10 -> 26.10, hâlâ açık
    assert ariza["G3"].value == 2      # 04.10 -> 06.10, çözüm süresi
    assert ariza["F2"].value is None


def test_bos_rapor_hata_vermez(tmp_path):
    bos = {"ozet": {}, "cihazlar": [], "bakimlar": [], "kalibrasyonlar": [], "arizalar": []}
    kitap = olustur(tmp_path, bos)

    assert kitap["Cihaz Özeti"].max_row == 1
    assert kitap["Bakım Geçmişi"].max_row == 1
    assert kitap["Arızalar"].max_row == 1
    assert satirlar(kitap["Özet"])["Rapordaki cihazlar"][0] == 0
    assert all(kayitlar == [] for kayitlar in tablolar(kitap["Durum Raporu"]).values())


@pytest.mark.parametrize("kotu", ["=HYPERLINK(\"http://x\",\"a\")", "+1+1", "-2+3", "@SUM(A1)"])
def test_formul_enjeksiyonu_engellenir(tmp_path, kotu):
    """Kullanıcının yazdığı '=' ile başlayan metin Excel'de formül olarak çalışmamalı."""
    bakimlar = [{**VERI["bakimlar"][0], "yapilan_islem": kotu}] + VERI["bakimlar"][1:]
    kitap = olustur(tmp_path, {**VERI, "bakimlar": bakimlar})

    durum = kitap["Durum Raporu"]
    hucreler = [kitap["Bakım Geçmişi"]["F2"], kitap["Cihaz Özeti"]["J2"]]
    hucreler += [
        durum.cell(r, 6) for r in range(1, durum.max_row + 1) if durum.cell(r, 6).value == kotu
    ]

    assert len(hucreler) >= 3   # Bakım Geçmişi, Cihaz Özeti ve Durum Raporu'ndaki tüm kopyalar
    for hucre in hucreler:
        assert hucre.value == kotu
        assert hucre.data_type == "s"