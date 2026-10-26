"""TKYS - Cihaz raporunu Excel (.xlsx) dosyasına yazar.

Ekrandan bağımsızdır; API'den gelen rapor verisini ve filtre özetini alır.
Rapor 6 sayfadan oluşur: Özet, Durum Raporu, Cihaz Özeti, Bakım Geçmişi,
Kalibrasyon, Arızalar.
"""

from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASLIK_DOLGU = PatternFill("solid", fgColor="E8F0FF")
BASLIK_FONT = Font(bold=True, color="16294B")
BOLUM_DOLGU = PatternFill("solid", fgColor="16294B")
BOLUM_FONT = Font(bold=True, color="FFFFFF")
ACIKLAMA_FONT = Font(italic=True, color="5A6B85")
UYARI_FONT = Font(bold=True, color="C0392B")
IYI_FONT = Font(bold=True, color="16803C")
DIKKAT_FONT = Font(bold=True, color="B86B08")

# Excel bu karakterlerle başlayan metni formül sayabilir.
FORMUL_BASLANGICLARI = ("=", "+", "-", "@", "\t", "\r")

TARIH_BICIMI = "dd.mm.yyyy"
TARIH_SAAT_BICIMI = "dd.mm.yyyy hh:mm"
PARA_BICIMI = '#,##0.00 "₺"'
GUN_BICIMI = '0.0 "gün"'

KALIBRASYON_UYARI_GUN = 30

DURUM_ADLARI = {
    "Kullanima_Hazir": "Kullanıma Hazır",
    "Rezerve": "Rezerve",
    "Kullanimda": "Kullanımda",
    "Arizali": "Arızalı",
    "Bakimda": "Bakımda",
    "Kalibrasyonda": "Kalibrasyonda",
    "Kullanim_Disi": "Kullanım Dışı",
}
DURUM_SIRASI = list(DURUM_ADLARI.values())

KAL_GECERLI = f"{KALIBRASYON_UYARI_GUN} günden fazla süresi var"
KAL_YAKINDA = f"{KALIBRASYON_UYARI_GUN} gün içinde dolacak"
KAL_DOLMUS = "Süresi dolmuş"
KAL_YOK = "Kalibrasyon kaydı yok"

BAKIM_DEVAM = "Devam ediyor"
BAKIM_PLANLI = "Planlanmış (henüz başlamadı)"
BAKIM_BITTI = "Tamamlandı"


# ---------------------------------------------------------------- yardımcılar
def _tarih_degeri(deger):
    """ISO metnini gerçek tarihe çevirir; boşsa None, çözülemezse metni döndürür."""
    if not deger:
        return None
    try:
        return datetime.fromisoformat(deger)
    except (TypeError, ValueError):
        return deger


def _gun(deger):
    """Tarih/tarih-saat değerini date'e çevirir; çevrilemezse None."""
    if isinstance(deger, datetime):
        return deger.date()
    if isinstance(deger, date):
        return deger
    return None


def _hucre_yaz(sayfa, satir, sutun, deger, bicim=None):
    hucre = sayfa.cell(row=satir, column=sutun)
    hucre.value = deger

    if isinstance(deger, str):
        # Formül enjeksiyonunu engelle: metin formül değil, düz metin olarak saklanır.
        hucre.data_type = "s"
        if deger.startswith(FORMUL_BASLANGICLARI):
            hucre.quotePrefix = True
    elif bicim:
        hucre.number_format = bicim

    return hucre


def _hizala(hucre, deger):
    hucre.alignment = Alignment(
        horizontal="left" if isinstance(deger, str) or deger is None else "center",
        vertical="top",
        wrap_text=isinstance(deger, str),
    )


def _sutun_genislikleri(sayfa, kolon_adlari, satirlar, ust_sinir=60):
    for sutun, kolon_adi in enumerate(kolon_adlari, start=1):
        uzunluklar = [len(str(kolon_adi))] + [
            len(str(satir[sutun - 1])) for satir in satirlar if satir[sutun - 1] is not None
        ]
        sayfa.column_dimensions[get_column_letter(sutun)].width = min(max(uzunluklar) + 3, ust_sinir)


def _sayfa_olustur(kitap, baslik, kolonlar, satirlar):
    """kolonlar: [(başlık, biçim)]; satirlar: değer listeleri."""
    sayfa = kitap.create_sheet(baslik)

    for sutun, (kolon_adi, _) in enumerate(kolonlar, start=1):
        hucre = sayfa.cell(row=1, column=sutun, value=kolon_adi)
        hucre.font = BASLIK_FONT
        hucre.fill = BASLIK_DOLGU
        hucre.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for satir_no, satir in enumerate(satirlar, start=2):
        for sutun, (deger, (_, bicim)) in enumerate(zip(satir, kolonlar), start=1):
            _hizala(_hucre_yaz(sayfa, satir_no, sutun, deger, bicim), deger)

    _sutun_genislikleri(sayfa, [k[0] for k in kolonlar], satirlar)
    sayfa.freeze_panes = "A2"
    if satirlar:
        sayfa.auto_filter.ref = sayfa.dimensions
    return sayfa


# ------------------------------------------------------------ iş kuralları
def _kalibrasyon_durumu(gecerlilik, bugun):
    gun = _gun(gecerlilik)
    if gun is None:
        return KAL_YOK
    fark = (gun - bugun).days
    if fark < 0:
        return KAL_DOLMUS
    if fark <= KALIBRASYON_UYARI_GUN:
        return KAL_YAKINDA
    return KAL_GECERLI


def _bakim_durumu(bakim, bugun):
    """Bakım başlangıç/bitiş tarihine göre bugünkü durumu."""
    bas = _gun(_tarih_degeri(bakim.get("bakim_baslangic_tarihi") or bakim.get("bakim_tarihi")))
    bit = _gun(_tarih_degeri(bakim.get("bakim_bitis_tarihi") or bakim.get("bakim_tarihi")))
    if bas is not None and bas > bugun:
        return BAKIM_PLANLI
    if bit is not None and bit >= bugun:
        return BAKIM_DEVAM
    return BAKIM_BITTI


def _gun_sayisi(bas, bit):
    bas, bit = _gun(bas), _gun(bit)
    if bas is None or bit is None:
        return None
    return max((bit - bas).days, 0)


def _ariza_suresi_gun(ariza, bugun):
    """Çözülmüşse çözüm süresi, açıksa açık kalma süresi (gün)."""
    bas = _gun(_tarih_degeri(ariza["bildirim_tarihi"]))
    if bas is None:
        return None
    bitis = _gun(_tarih_degeri(ariza["cozum_tarihi"])) if ariza.get("cozum_tarihi") else bugun
    if bitis is None:
        return None
    return max((bitis - bas).days, 0)


def _cihaz_ozetlerini_hesapla(veri, bugun):
    """Her cihaz için bakım/kalibrasyon/arıza özetini çıkarır."""
    ozetler = []

    for cihaz in veri.get("cihazlar", []):
        cid = cihaz["cihaz_id"]
        bakimlar = [b for b in veri.get("bakimlar", []) if b.get("cihaz_id") == cid]
        kalibrasyonlar = [k for k in veri.get("kalibrasyonlar", []) if k.get("cihaz_id") == cid]
        arizalar = [a for a in veri.get("arizalar", []) if a.get("cihaz_id") == cid]

        maliyetler = [b.get("maliyet") or 0 for b in bakimlar]
        son_bakim = bakimlar[0] if bakimlar else None  # API en yeniden eskiye sıralar

        son_kal = cihaz.get("son_kalibrasyon") or {}
        son_gecerlilik = _tarih_degeri(son_kal.get("gecerlilik_tarihi"))
        if _gun(son_gecerlilik) is None:
            son_gecerlilik = None

        acik_arizalar = [a for a in arizalar if a["durum"] == "Açık"]
        son_ariza = (acik_arizalar or arizalar or [None])[0]

        sonraki_bakim = _tarih_degeri(cihaz.get("sonraki_bakim_tarihi"))
        if sonraki_bakim is None and son_bakim:
            sonraki_bakim = _tarih_degeri(son_bakim["sonraki_bakim_tarihi"])

        ozetler.append({
            "cihaz": cihaz,
            "durum": DURUM_ADLARI.get(cihaz["durum"], cihaz["durum"]),
            "bakim_sayisi": len(bakimlar),
            "devam_eden_bakim": sum(1 for b in bakimlar if _bakim_durumu(b, bugun) == BAKIM_DEVAM),
            "toplam_maliyet": sum(maliyetler),
            "ortalama_maliyet": (sum(maliyetler) / len(maliyetler)) if maliyetler else None,
            "son_bakim_tarihi": _tarih_degeri(son_bakim["bakim_tarihi"]) if son_bakim else None,
            "son_bakim_islemi": son_bakim.get("yapilan_islem") if son_bakim else None,
            "sonraki_bakim": sonraki_bakim,
            "kalibrasyon_sayisi": len(kalibrasyonlar),
            "kalibrasyon_gecerlilik": son_gecerlilik,
            "kalibrasyon_durumu": _kalibrasyon_durumu(son_gecerlilik, bugun),
            "kalibrasyon": son_kal,
            "ariza_sayisi": len(arizalar),
            "acik_ariza": len(acik_arizalar),
            "son_ariza": son_ariza.get("ariza_aciklamasi") if son_ariza else None,
        })

    return ozetler


# ------------------------------------------------------------- Özet sayfası
def _kodlar(liste):
    return ", ".join(liste) if liste else "Yok"


def _ozet_sayfasini_yaz(sayfa, veri, filtre_ozeti, ozetler, bugun):
    sayfa.sheet_view.showGridLines = False
    satir = 1

    def yaz(a, b=None, c=None, bicim=None, kalin=False):
        nonlocal satir
        h1 = sayfa.cell(row=satir, column=1, value=a)
        h1.font = BASLIK_FONT if kalin else Font()
        h1.alignment = Alignment(vertical="top", wrap_text=True)
        if b is not None:
            h2 = _hucre_yaz(sayfa, satir, 2, b, bicim)
            h2.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if c is not None:
            h3 = _hucre_yaz(sayfa, satir, 3, c)
            h3.font = ACIKLAMA_FONT
            h3.alignment = Alignment(vertical="top", wrap_text=True)
        satir += 1

    def bolum(metin):
        nonlocal satir
        satir += 1
        for sutun in (1, 2, 3):
            hucre = sayfa.cell(row=satir, column=sutun, value=metin if sutun == 1 else None)
            hucre.fill = BOLUM_DOLGU
            hucre.font = BOLUM_FONT
        satir += 1

    baslik = sayfa.cell(row=satir, column=1, value="CİHAZ RAPORU — Bakım, Kalibrasyon ve Arıza")
    baslik.font = Font(bold=True, size=15, color="16294B")
    satir += 1
    alt = sayfa.cell(
        row=satir, column=1,
        value="Bu rapor, sisteme girilmiş cihaz bakım, kalibrasyon ve arıza kayıtlarından oluşturulmuştur. "
              "Kayıt kayıt ayrıntı için 'Durum Raporu' sayfasına bakın.",
    )
    alt.font = ACIKLAMA_FONT
    satir += 1

    # ---- Rapor bilgileri ----
    bolum("RAPOR BİLGİLERİ")
    yaz("Rapor Tarihi", datetime.now(), "Raporun Excel'e aktarıldığı an.", TARIH_SAAT_BICIMI, kalin=True)

    for ad, deger in filtre_ozeti:
        if ad == "Dönem":
            aciklama = (
                "Tarih sınırı uygulanmadı; kayıtlı tüm dönemler rapora dahil." if deger == "Tümü"
                else "Bu tarihler arasındaki kayıtlar rapora dahil edildi (arızada bildirim "
                     "ya da çözüm tarihi bu aralıkta olanlar)."
            )
        elif ad == "Cihaz":
            aciklama = (
                "Belirli bir cihaz seçilmedi; envanterdeki tüm cihazlar rapora dahil."
                if deger == "Tüm Cihazlar"
                else "Rapor yalnızca bu cihaza aittir."
            )
        else:
            aciklama = None
        yaz(ad, deger, aciklama, kalin=True)

    yaz("Rapordaki cihazlar", len(ozetler),
        ", ".join(f"{o['cihaz']['envanter_kodu']} ({o['cihaz']['cihaz_adi']})" for o in ozetler) or "Cihaz bulunamadı.",
        kalin=True)

    # ---- Bakım ----
    bakimlar = veri.get("bakimlar", [])
    ozet = veri.get("ozet", {})
    siniflar = {BAKIM_DEVAM: [], BAKIM_PLANLI: [], BAKIM_BITTI: []}
    for b in bakimlar:
        siniflar[_bakim_durumu(b, bugun)].append(b)
    cihaz_kodu = lambda kayitlar: _kodlar(sorted({k["cihaz"] for k in kayitlar}))

    bolum("BAKIM")
    yaz("Bakım kaydı (toplam)", ozet.get("bakim_sayisi", len(bakimlar)),
        "Seçilen dönemde girilen tüm bakım kayıtları.", kalin=True)
    yaz("Devam eden bakım", len(siniflar[BAKIM_DEVAM]),
        f"Şu an süren bakımlar: {cihaz_kodu(siniflar[BAKIM_DEVAM])}", kalin=True)
    yaz("Planlanmış bakım", len(siniflar[BAKIM_PLANLI]),
        f"Başlangıç tarihi henüz gelmemiş bakımlar: {cihaz_kodu(siniflar[BAKIM_PLANLI])}", kalin=True)
    yaz("Tamamlanan bakım", len(siniflar[BAKIM_BITTI]),
        f"Bitiş tarihi geçmiş bakımlar: {cihaz_kodu(siniflar[BAKIM_BITTI])}", kalin=True)
    toplam_maliyet = ozet.get("toplam_bakim_maliyeti", 0)
    yaz("Toplam bakım maliyeti", toplam_maliyet, "Bakım kayıtlarındaki maliyetlerin toplamı.", PARA_BICIMI, kalin=True)
    if bakimlar:
        yaz("Ortalama bakım maliyeti", toplam_maliyet / len(bakimlar),
            "Toplam maliyetin bakım kaydı sayısına bölümü.", PARA_BICIMI, kalin=True)
        en_pahali = max(bakimlar, key=lambda b: b.get("maliyet") or 0)
        yaz("En yüksek maliyetli bakım", en_pahali.get("maliyet") or 0,
            f"{en_pahali['cihaz']} — {en_pahali.get('yapilan_islem') or 'işlem belirtilmemiş'}",
            PARA_BICIMI, kalin=True)

    # ---- Kalibrasyon (güncel durum) ----
    dolmus = [o for o in ozetler if o["kalibrasyon_durumu"] == KAL_DOLMUS]
    yakinda = [o for o in ozetler if o["kalibrasyon_durumu"] == KAL_YAKINDA]
    gecerli = [o for o in ozetler if o["kalibrasyon_durumu"] == KAL_GECERLI]
    kayitsiz = [o for o in ozetler if o["kalibrasyon_durumu"] == KAL_YOK]
    kodlar = lambda liste: _kodlar([o["cihaz"]["envanter_kodu"] for o in liste])

    bolum("KALİBRASYON (bugünkü durum)")
    yaz("Kalibrasyon kaydı (dönemde)", ozet.get("kalibrasyon_sayisi", len(veri.get("kalibrasyonlar", []))),
        "Seçilen dönemde girilen kalibrasyon kaydı sayısı.", kalin=True)
    yaz("Süresi dolmuş", len(dolmus), f"Geçerlilik tarihi geçmiş cihazlar: {kodlar(dolmus)}", kalin=True)
    yaz(f"{KALIBRASYON_UYARI_GUN} gün içinde dolacak", len(yakinda),
        f"Geçerliliği {KALIBRASYON_UYARI_GUN} günden az kalmış cihazlar: {kodlar(yakinda)}", kalin=True)
    yaz(f"{KALIBRASYON_UYARI_GUN} günden fazla süresi var", len(gecerli),
        f"Kalibrasyonu geçerli ve süresi rahat cihazlar: {kodlar(gecerli)}", kalin=True)
    yaz("Kalibrasyon kaydı yok", len(kayitsiz), f"Hiç kalibrasyon kaydı bulunmayan cihazlar: {kodlar(kayitsiz)}", kalin=True)
    yaz("Not", None, "Kalibrasyon durumu bugüne göre hesaplanır; tarih filtresinden etkilenmez.")

    # ---- Arıza ----
    arizalar = veri.get("arizalar", [])
    acik = [a for a in arizalar if a["durum"] == "Açık"]
    cozulen = [a for a in arizalar if a["durum"] != "Açık"]

    bolum("ARIZA")
    yaz("Arıza kaydı (toplam)", ozet.get("ariza_sayisi", len(arizalar)),
        "Seçilen dönemde bildirilen ya da çözülen arızalar.", kalin=True)
    yaz("Çözümü süren arıza", len(acik),
        f"Henüz çözülmemiş, çözüm için uğraşılan arızalar: {cihaz_kodu(acik)}", kalin=True)
    yaz("Çözülen arıza", len(cozulen), f"Çözüm tarihi girilmiş arızalar: {cihaz_kodu(cozulen)}", kalin=True)
    sureler = [s for s in (_ariza_suresi_gun(a, bugun) for a in cozulen) if s is not None]
    if sureler:
        yaz("Ortalama çözüm süresi", sum(sureler) / len(sureler),
            "Bildirimden çözüme kadar geçen ortalama süre.", GUN_BICIMI, kalin=True)
    if acik:
        en_eski = min(acik, key=lambda a: a["bildirim_tarihi"] or "")
        yaz("En uzun süredir açık arıza", _ariza_suresi_gun(en_eski, bugun),
            f"{en_eski['cihaz']} — {en_eski.get('ariza_aciklamasi') or ''}", '0 "gün"', kalin=True)

    # ---- Durum dağılımı ----
    bolum("CİHAZLARIN GÜNCEL DURUMU")
    yaz("Durum", "Cihaz Sayısı", "Bu durumdaki cihazlar", kalin=True)
    for durum in DURUM_SIRASI:
        liste = [o for o in ozetler if o["durum"] == durum]
        if not liste:
            continue
        yaz(durum, len(liste),
            ", ".join(f"{o['cihaz']['envanter_kodu']} ({o['cihaz']['cihaz_adi']})" for o in liste),
            kalin=True)
    diger = [o for o in ozetler if o["durum"] not in DURUM_SIRASI]
    if diger:
        yaz("Diğer", len(diger), kodlar(diger), kalin=True)
    yaz("Not", None, "Durum, raporun alındığı andaki güncel durumdur; tarih filtresinden etkilenmez.")

    # ---- Sayfa rehberi ----
    bolum("SAYFA REHBERİ")
    for ad, aciklama in [
        ("Durum Raporu", "Devam eden/tamamlanan bakımlar, çözülen/süren arızalar ve kalibrasyon durumları, kayıt kayıt."),
        ("Cihaz Özeti", "Her cihaz için tek satır: güncel durum, bakım maliyeti, kalibrasyon durumu, açık arıza."),
        ("Bakım Geçmişi", "Tüm bakım kayıtları: kim, ne yaptı, ne kadar tuttu, bakım durumu."),
        ("Kalibrasyon", "Seçilen dönemdeki kalibrasyon kayıtları ve her birinin geçerlilik durumu."),
        ("Arızalar", "Arıza kayıtları, durumları ve çözüm (ya da açık kalma) süreleri."),
    ]:
        yaz(ad, None, aciklama, kalin=True)

    sayfa.column_dimensions["A"].width = 34
    sayfa.column_dimensions["B"].width = 26
    sayfa.column_dimensions["C"].width = 90


# ----------------------------------------------------------- Durum Raporu
DURUM_SUTUNLARI = [26, 18, 18, 18, 22, 48, 16]


def _durum_tablosu(sayfa, satir, baslik, aciklama, kolonlar, satirlar, vurgu_sutunu=None, vurgu_font=None):
    """Başlık, açıklama ve tabloyu yazar; bir sonraki boş satırı döndürür.

    kolonlar: [(başlık, biçim)] en fazla 7 sütun; satirlar: değer listeleri.
    """
    hucre = sayfa.cell(row=satir, column=1, value=f"{baslik}  ({len(satirlar)} kayıt)")
    hucre.font = BOLUM_FONT
    for sutun in range(1, len(DURUM_SUTUNLARI) + 1):
        sayfa.cell(row=satir, column=sutun).fill = BOLUM_DOLGU
    satir += 1

    ac = sayfa.cell(row=satir, column=1, value=aciklama)
    ac.font = ACIKLAMA_FONT
    satir += 1

    if not satirlar:
        bos = sayfa.cell(row=satir, column=1, value="Bu kategoride kayıt yok.")
        bos.font = ACIKLAMA_FONT
        return satir + 2

    for sutun, (kolon_adi, _) in enumerate(kolonlar, start=1):
        h = sayfa.cell(row=satir, column=sutun, value=kolon_adi)
        h.font = BASLIK_FONT
        h.fill = BASLIK_DOLGU
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    satir += 1

    for deger_listesi in satirlar:
        for sutun, (deger, (_, bicim)) in enumerate(zip(deger_listesi, kolonlar), start=1):
            h = _hucre_yaz(sayfa, satir, sutun, deger, bicim)
            _hizala(h, deger)
            if vurgu_sutunu == sutun and vurgu_font is not None:
                h.font = vurgu_font
        satir += 1

    return satir + 1


def _durum_raporunu_yaz(sayfa, veri, ozetler, bugun):
    sayfa.sheet_view.showGridLines = False
    for sutun, genislik in enumerate(DURUM_SUTUNLARI, start=1):
        sayfa.column_dimensions[get_column_letter(sutun)].width = genislik

    baslik = sayfa.cell(row=1, column=1, value="DURUM RAPORU — Kategorilere göre kayıt kayıt")
    baslik.font = Font(bold=True, size=15, color="16294B")
    sayfa.cell(row=2, column=1, value=f"Durumlar {bugun.strftime('%d.%m.%Y')} tarihine göre hesaplanmıştır.").font = ACIKLAMA_FONT
    satir = 4

    # ---------------- BAKIM ----------------
    bakimlar = veri.get("bakimlar", [])
    siniflar = {BAKIM_DEVAM: [], BAKIM_PLANLI: [], BAKIM_BITTI: []}
    for b in bakimlar:
        siniflar[_bakim_durumu(b, bugun)].append(b)

    def bakim_satiri(b, son_sutun):
        bas = _tarih_degeri(b.get("bakim_baslangic_tarihi") or b["bakim_tarihi"])
        bit = _tarih_degeri(b.get("bakim_bitis_tarihi") or b["bakim_tarihi"])
        return [
            b["cihaz"], bas, bit, son_sutun(bas, bit),
            b.get("bakim_yapan_kisi"), b.get("yapilan_islem"), b.get("maliyet", 0),
        ]

    bakim_kolonlari = lambda ucuncu: [
        ("Cihaz", None), ("Bakım Başlangıcı", TARIH_BICIMI), ("Bakım Bitişi", TARIH_BICIMI),
        (ucuncu, None), ("Yapan Kişi", None), ("Yapılan İşlem", None), ("Maliyet", PARA_BICIMI),
    ]

    satir = _durum_tablosu(
        sayfa, satir, "BAKIM — Devam eden",
        "Bakımı şu an süren cihazlar (başlangıç tarihi gelmiş, bitiş tarihi henüz geçmemiş).",
        bakim_kolonlari("Kaç Gündür Sürüyor"),
        [bakim_satiri(b, lambda bas, bit: _gun_sayisi(bas, bugun)) for b in siniflar[BAKIM_DEVAM]],
    )
    satir = _durum_tablosu(
        sayfa, satir, "BAKIM — Planlanmış",
        "Başlangıç tarihi henüz gelmemiş, ileri tarihe planlanmış bakımlar.",
        bakim_kolonlari("Başlamasına Kalan Gün"),
        [bakim_satiri(b, lambda bas, bit: _gun_sayisi(bugun, bas)) for b in siniflar[BAKIM_PLANLI]],
    )
    satir = _durum_tablosu(
        sayfa, satir, "BAKIM — Tamamlanan",
        "Bitiş tarihi geçmiş, tamamlanmış bakımlar.",
        bakim_kolonlari("Süre (gün)"),
        [bakim_satiri(b, _gun_sayisi) for b in siniflar[BAKIM_BITTI]],
    )

    # Planlı sonraki bakımlar (güncel)
    planli = []
    for o in ozetler:
        sonraki = o["sonraki_bakim"]
        gun = _gun(sonraki)
        if gun is None:
            continue
        kalan = (gun - bugun).days
        if kalan < 0:
            durum = f"{-kalan} gün gecikmiş"
        elif kalan <= KALIBRASYON_UYARI_GUN:
            durum = f"{KALIBRASYON_UYARI_GUN} gün içinde"
        else:
            durum = "Zamanı var"
        planli.append((kalan, [
            f"{o['cihaz']['envanter_kodu']} - {o['cihaz']['cihaz_adi']}", sonraki, kalan, durum,
            None, o["son_bakim_islemi"], None,
        ]))
    planli.sort(key=lambda x: x[0])
    satir = _durum_tablosu(
        sayfa, satir, "BAKIM — Bir sonraki planlı bakım",
        "Her cihazın kayıtlı sonraki bakım tarihi (negatif gün: bakımı gecikmiş). Tarih filtresinden etkilenmez.",
        [("Cihaz", None), ("Sonraki Bakım", TARIH_BICIMI), ("Kalan Gün", None), ("Durum", None),
         ("", None), ("Son Yapılan Bakım İşlemi", None), ("", None)],
        [p[1] for p in planli], vurgu_sutunu=4, vurgu_font=DIKKAT_FONT,
    )

    # ---------------- ARIZA ----------------
    arizalar = veri.get("arizalar", [])
    acik = [a for a in arizalar if a["durum"] == "Açık"]
    cozulen = [a for a in arizalar if a["durum"] != "Açık"]
    acik.sort(key=lambda a: a["bildirim_tarihi"] or "")
    cozulen.sort(key=lambda a: a["cozum_tarihi"] or "", reverse=True)

    satir = _durum_tablosu(
        sayfa, satir, "ARIZA — Çözümü süren",
        "Henüz çözülmemiş, çözüm için uğraşılan arızalar (en uzun süredir açık olan en üstte).",
        [("Cihaz", None), ("Bildirim Tarihi", TARIH_SAAT_BICIMI), ("Kaç Gündür Açık", None),
         ("Bildiren", None), ("", None), ("Arıza Açıklaması", None), ("", None)],
        [[a["cihaz"], _tarih_degeri(a["bildirim_tarihi"]), _ariza_suresi_gun(a, bugun),
          a.get("bildiren_personel"), None, a.get("ariza_aciklamasi"), None] for a in acik],
        vurgu_sutunu=3, vurgu_font=UYARI_FONT,
    )
    satir = _durum_tablosu(
        sayfa, satir, "ARIZA — Çözülen",
        "Çözüm tarihi girilmiş arızalar (en son çözülen en üstte).",
        [("Cihaz", None), ("Bildirim Tarihi", TARIH_SAAT_BICIMI), ("Çözüm Tarihi", TARIH_SAAT_BICIMI),
         ("Çözüm Süresi (gün)", None), ("Bildiren", None), ("Arıza Açıklaması", None), ("", None)],
        [[a["cihaz"], _tarih_degeri(a["bildirim_tarihi"]), _tarih_degeri(a["cozum_tarihi"]),
          _ariza_suresi_gun(a, bugun), a.get("bildiren_personel"), a.get("ariza_aciklamasi"), None]
         for a in cozulen],
    )

    # ---------------- KALİBRASYON ----------------
    def kal_satiri(o, kalan_degeri):
        cihaz = o["cihaz"]
        kal = o["kalibrasyon"]
        return [
            f"{cihaz['envanter_kodu']} - {cihaz['cihaz_adi']}",
            o["kalibrasyon_gecerlilik"], kalan_degeri,
            _tarih_degeri(kal.get("islem_tarihi")), kal.get("sertifika_no"),
            kal.get("kalibrasyon_yapan_kurum"), None,
        ]

    def kal_kolonlari(ucuncu):
        return [("Cihaz", None), ("Geçerlilik Bitişi", TARIH_BICIMI), (ucuncu, None),
                ("Son Kalibrasyon İşlemi", TARIH_SAAT_BICIMI), ("Sertifika No", None),
                ("Kalibrasyonu Yapan Kurum", None), ("", None)]

    def kal_gunu(o):
        return (_gun(o["kalibrasyon_gecerlilik"]) - bugun).days

    kal_gruplari = [
        (KAL_DOLMUS, "KALİBRASYON — Süresi dolmuş",
         "Son kalibrasyonunun geçerlilik tarihi geçmiş cihazlar. Yeniden kalibrasyon gerekir.",
         "Kaç Gün Önce Doldu", lambda o: -kal_gunu(o), UYARI_FONT),
        (KAL_YAKINDA, f"KALİBRASYON — {KALIBRASYON_UYARI_GUN} gün içinde dolacak",
         f"Geçerliliği {KALIBRASYON_UYARI_GUN} günden az kalmış cihazlar. Yakında yeniden kalibre edilmeli.",
         "Kalan Gün", kal_gunu, DIKKAT_FONT),
        (KAL_GECERLI, f"KALİBRASYON — {KALIBRASYON_UYARI_GUN} günden fazla süresi var",
         "Kalibrasyonu geçerli ve süresi rahat cihazlar.",
         "Kalan Gün", kal_gunu, IYI_FONT),
    ]
    for durum, baslik_metni, aciklama, kolon, hesapla, font in kal_gruplari:
        grup = sorted((o for o in ozetler if o["kalibrasyon_durumu"] == durum), key=kal_gunu)
        satir = _durum_tablosu(
            sayfa, satir, baslik_metni, aciklama, kal_kolonlari(kolon),
            [kal_satiri(o, hesapla(o)) for o in grup], vurgu_sutunu=3, vurgu_font=font,
        )

    kayitsiz = [o for o in ozetler if o["kalibrasyon_durumu"] == KAL_YOK]
    satir = _durum_tablosu(
        sayfa, satir, "KALİBRASYON — Kaydı olmayan",
        "Sistemde hiç kalibrasyon kaydı bulunmayan cihazlar.",
        [("Cihaz", None), ("Güncel Durum", None), ("", None), ("", None), ("", None), ("", None), ("", None)],
        [[f"{o['cihaz']['envanter_kodu']} - {o['cihaz']['cihaz_adi']}", o["durum"], None, None, None, None, None]
         for o in kayitsiz],
    )


# ----------------------------------------------------------------- ana akış
def excel_dosyasi_olustur(veri, filtre_ozeti, dosya_yolu, bugun=None):
    """Rapor verisini 6 sayfalı .xlsx dosyası olarak kaydeder.

    veri: GET /api/cihaz-raporlari yanıtı
    filtre_ozeti: [(ad, değer)] biçiminde ekrandaki filtreler
    bugun: durum hesapları için tarih (varsayılan: bugün)
    """
    bugun = bugun or date.today()
    ozetler = _cihaz_ozetlerini_hesapla(veri, bugun)

    kitap = Workbook()
    ozet_sayfasi = kitap.active
    ozet_sayfasi.title = "Özet"
    _ozet_sayfasini_yaz(ozet_sayfasi, veri, filtre_ozeti, ozetler, bugun)

    _durum_raporunu_yaz(kitap.create_sheet("Durum Raporu"), veri, ozetler, bugun)

    _sayfa_olustur(
        kitap,
        "Cihaz Özeti",
        [("Envanter Kodu", None), ("Cihaz Adı", None), ("Marka / Model", None),
         ("Güncel Durum", None), ("Bakım Sayısı (dönemde)", None), ("Devam Eden Bakım", None),
         ("Toplam Bakım Maliyeti", PARA_BICIMI), ("Ortalama Bakım Maliyeti", PARA_BICIMI),
         ("Son Bakım Tarihi", TARIH_BICIMI), ("Son Bakım İşlemi", None),
         ("Sonraki Planlı Bakım", TARIH_BICIMI),
         ("Kalibrasyon Kaydı (dönemde)", None), ("Kalibrasyon Geçerlilik Bitişi", TARIH_BICIMI),
         ("Kalibrasyon Durumu", None), ("Arıza Sayısı (dönemde)", None), ("Çözümü Süren Arıza", None),
         ("Son / Açık Arıza Açıklaması", None)],
        [
            [
                o["cihaz"]["envanter_kodu"], o["cihaz"]["cihaz_adi"], o["cihaz"].get("marka_model"),
                o["durum"], o["bakim_sayisi"], o["devam_eden_bakim"],
                o["toplam_maliyet"], o["ortalama_maliyet"],
                o["son_bakim_tarihi"], o["son_bakim_islemi"],
                o["sonraki_bakim"],
                o["kalibrasyon_sayisi"], o["kalibrasyon_gecerlilik"],
                o["kalibrasyon_durumu"], o["ariza_sayisi"], o["acik_ariza"],
                o["son_ariza"],
            ]
            for o in ozetler
        ],
    )

    _sayfa_olustur(
        kitap,
        "Bakım Geçmişi",
        [("Cihaz", None), ("Bakım Başlangıcı", TARIH_BICIMI), ("Bakım Bitişi", TARIH_BICIMI),
         ("Bakım Durumu", None), ("Yapan Kişi", None), ("Yapılan İşlem", None),
         ("Maliyet", PARA_BICIMI), ("Sonraki Bakım", TARIH_BICIMI)],
        [
            [
                b["cihaz"],
                _tarih_degeri(b.get("bakim_baslangic_tarihi") or b["bakim_tarihi"]),
                _tarih_degeri(b.get("bakim_bitis_tarihi") or b["bakim_tarihi"]),
                _bakim_durumu(b, bugun),
                b.get("bakim_yapan_kisi"),
                b.get("yapilan_islem"),
                b.get("maliyet", 0),
                _tarih_degeri(b["sonraki_bakim_tarihi"]),
            ]
            for b in veri.get("bakimlar", [])
        ],
    )

    _sayfa_olustur(
        kitap,
        "Kalibrasyon",
        [("Cihaz", None), ("İşlem Tarihi", TARIH_SAAT_BICIMI), ("Kurum", None),
         ("Sertifika No", None), ("Geçerlilik Tarihi", TARIH_BICIMI), ("Geçerlilik Durumu", None)],
        [
            [
                k["cihaz"],
                _tarih_degeri(k["islem_tarihi"]),
                k.get("kalibrasyon_yapan_kurum"),
                k.get("sertifika_no"),
                _tarih_degeri(k["gecerlilik_tarihi"]),
                _kalibrasyon_durumu(_tarih_degeri(k["gecerlilik_tarihi"]), bugun),
            ]
            for k in veri.get("kalibrasyonlar", [])
        ],
    )

    _sayfa_olustur(
        kitap,
        "Arızalar",
        [("Cihaz", None), ("Bildirim Tarihi", TARIH_SAAT_BICIMI), ("Bildiren", None),
         ("Açıklama", None), ("Durum", None), ("Çözüm Tarihi", TARIH_SAAT_BICIMI),
         ("Çözüm / Açık Kalma Süresi (gün)", None)],
        [
            [
                a["cihaz"],
                _tarih_degeri(a["bildirim_tarihi"]),
                a.get("bildiren_personel"),
                a.get("ariza_aciklamasi"),
                "Çözümü sürüyor" if a["durum"] == "Açık" else a["durum"],
                _tarih_degeri(a["cozum_tarihi"]),
                _ariza_suresi_gun(a, bugun),
            ]
            for a in veri.get("arizalar", [])
        ],
    )

    kitap.save(dosya_yolu)